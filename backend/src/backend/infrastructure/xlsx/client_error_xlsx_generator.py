from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from backend.infrastructure.xlsx.client_xlsx_parser import RejectedClientRow

HEADERS = (
    "ПІБ",
    "Телефон 1",
    "Телефон 2",
    "Вулиця 1",
    "Будинок 1",
    "Квартира 1",
    "Під'їзд 1",
    "Поверх 1",
    "Домофон 1",
    "Район 1",
    "Коментар 1",
    "Вулиця 2",
    "Будинок 2",
    "Квартира 2",
    "Під'їзд 2",
    "Поверх 2",
    "Домофон 2",
    "Район 2",
    "Коментар 2",
    "Помилка",
)

BOLD = Font(bold=True)


class ClientErrorXlsxGenerator:
    def generate(
        self,
        rows: list[RejectedClientRow],
        headers: list[str] | None = None,
    ) -> bytes:
        wb = Workbook()
        ws = wb.active
        if ws is None:
            wb.close()
            return b""

        header_row = [*headers, "Помилка"] if headers is not None else HEADERS

        for col_idx, header in enumerate(header_row, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = BOLD

        for row_offset, rejected in enumerate(rows, start=2):
            for col_idx, value in enumerate(rejected.raw_values, start=1):
                ws.cell(row=row_offset, column=col_idx, value=value)
            ws.cell(
                row=row_offset,
                column=len(header_row),
                value=rejected.error_message,
            )

        buf = BytesIO()
        wb.save(buf)
        wb.close()
        return buf.getvalue()
