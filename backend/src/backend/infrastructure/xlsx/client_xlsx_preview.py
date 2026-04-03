from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

MAX_PREVIEW_COLUMNS = 50
SAMPLE_ROWS_COUNT = 3


@dataclass(frozen=True)
class ColumnPreview:
    index: int
    header: str | None
    sample_values: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class PreviewResult:
    columns: list[ColumnPreview]
    total_rows: int


def preview_xlsx(file_bytes: bytes) -> PreviewResult:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        return PreviewResult(columns=[], total_rows=0)

    rows: list[list[str | None]] = []
    total_rows = 0
    max_col = 0

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [
            str(cell).strip() if cell is not None else None
            for cell in row[:MAX_PREVIEW_COLUMNS]
        ]

        col_count = len(values)
        max_col = max(max_col, col_count)

        if row_idx <= 1 + SAMPLE_ROWS_COUNT:
            rows.append(values)

        if any(v is not None for v in values):
            total_rows = row_idx

    wb.close()

    if not rows:
        return PreviewResult(columns=[], total_rows=0)

    header_row = rows[0] if rows else []
    sample_rows = rows[1 : 1 + SAMPLE_ROWS_COUNT]

    columns: list[ColumnPreview] = []
    for col_idx in range(max_col):
        header = header_row[col_idx] if col_idx < len(header_row) else None
        samples: list[str] = []
        for sample_row in sample_rows:
            val = sample_row[col_idx] if col_idx < len(sample_row) else None
            if val is not None:
                samples.append(val)

        if header is not None or samples:
            columns.append(
                ColumnPreview(
                    index=col_idx + 1,
                    header=header,
                    sample_values=samples,
                )
            )

    data_rows = total_rows - 1 if total_rows > 0 else 0

    return PreviewResult(columns=columns, total_rows=data_rows)
