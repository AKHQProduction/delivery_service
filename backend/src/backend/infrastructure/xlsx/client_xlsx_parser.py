import logging
from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell

from backend.infrastructure.xlsx.column_mapping import (
    ColumnMapping,
    SystemField,
)

logger = logging.getLogger(__name__)

DATA_START_ROW = 4
MAX_COL = 19

COL_FULL_NAME = 1
COL_PHONE1 = 2
COL_PHONE2 = 3
COL_ADDR1_STREET = 4
COL_ADDR1_HOUSE = 5
COL_ADDR1_APARTMENT = 6
COL_ADDR1_ENTRANCE = 7
COL_ADDR1_FLOOR = 8
COL_ADDR1_INTERCOM = 9
COL_ADDR1_DISTRICT = 10
COL_ADDR1_COMMENT = 11
COL_ADDR2_STREET = 12
COL_ADDR2_HOUSE = 13
COL_ADDR2_APARTMENT = 14
COL_ADDR2_ENTRANCE = 15
COL_ADDR2_FLOOR = 16
COL_ADDR2_INTERCOM = 17
COL_ADDR2_DISTRICT = 18
COL_ADDR2_COMMENT = 19

RawCellValues = tuple[str | None, ...]


class ImportErrorKind(StrEnum):
    EMPTY_FIELD = "Пусте поле"
    DUPLICATE_IN_DB = "Дубль (у базі)"
    DUPLICATE_IN_FILE = "Дубль у файлі"
    INVALID_PHONE = "Невалідний телефон"


@dataclass(frozen=True)
class ParsedAddress:
    street: str
    house: str
    apartment: str | None
    entrance: str | None
    floor: str | None
    intercom: str | None
    district_name: str | None
    comment: str | None


@dataclass(frozen=True)
class ParsedClientRow:
    row_number: int
    full_name: str
    phone1: str
    phone2: str | None
    address1: ParsedAddress
    address2: ParsedAddress | None


@dataclass(frozen=True)
class RejectedClientRow:
    row_number: int
    raw_values: RawCellValues
    error_kind: ImportErrorKind
    error_detail: str | None = None

    @property
    def error_message(self) -> str:
        if self.error_detail:
            return f"{self.error_kind.value} ({self.error_detail})"
        return self.error_kind.value


@dataclass(frozen=True)
class ParseResult:
    valid: list[ParsedClientRow]
    rejected: list[RejectedClientRow]
    raw_by_row: dict[int, RawCellValues]


class ClientXlsxParser:
    def parse(
        self,
        file_bytes: bytes,
        column_mapping: ColumnMapping | None = None,
        data_start_row: int = DATA_START_ROW,
    ) -> ParseResult:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            wb.close()
            return ParseResult(valid=[], rejected=[], raw_by_row={})

        if column_mapping is not None:
            field_to_col: dict[SystemField, int] | None = {
                f: c for c, f in column_mapping.items()
            }
            max_col = max(column_mapping.keys())
        else:
            field_to_col = None
            max_col = MAX_COL

        valid: list[ParsedClientRow] = []
        rejected: list[RejectedClientRow] = []
        raw_by_row: dict[int, RawCellValues] = {}

        for row_idx, row in enumerate(
            ws.iter_rows(min_row=data_start_row, values_only=False),
            start=data_start_row,
        ):
            cells: dict[int, object] = {}
            for cell in row:
                if not isinstance(cell, EmptyCell) and cell.column is not None:
                    cells[cell.column] = cell.value

            raw = tuple(_str(cells.get(col)) for col in range(1, max_col + 1))
            raw_by_row[row_idx] = raw

            parsed = _parse_row(cells, field_to_col, row_idx, raw)
            if isinstance(parsed, RejectedClientRow):
                rejected.append(parsed)
            elif parsed is not None:
                valid.append(parsed)

        wb.close()
        logger.info(
            "Parsed XLSX: %d valid, %d rejected rows",
            len(valid),
            len(rejected),
        )
        return ParseResult(
            valid=valid, rejected=rejected, raw_by_row=raw_by_row
        )


def _parse_row(
    cells: dict[int, object],
    field_to_col: dict[SystemField, int] | None,
    row_idx: int,
    raw: RawCellValues,
) -> ParsedClientRow | RejectedClientRow | None:
    full_name = _get(cells, field_to_col, SystemField.FULL_NAME, COL_FULL_NAME)
    phone1 = _get(cells, field_to_col, SystemField.PHONE1, COL_PHONE1)
    street1 = _get(
        cells, field_to_col, SystemField.ADDR1_STREET, COL_ADDR1_STREET
    )
    house1 = _get(
        cells, field_to_col, SystemField.ADDR1_HOUSE, COL_ADDR1_HOUSE
    )

    if not full_name or not phone1 or not street1 or not house1:
        if any(v is not None for v in raw):
            return RejectedClientRow(
                row_number=row_idx,
                raw_values=raw,
                error_kind=ImportErrorKind.EMPTY_FIELD,
            )
        return None

    # address1 is guaranteed non-None since street1/house1 are validated above
    address1 = _build_address(
        cells, field_to_col, street1, house1, addr_prefix=1
    )
    address2 = _build_optional_address(cells, field_to_col, addr_prefix=2)

    return ParsedClientRow(
        row_number=row_idx,
        full_name=full_name,
        phone1=phone1,
        phone2=_get(cells, field_to_col, SystemField.PHONE2, COL_PHONE2),
        address1=address1,
        address2=address2,
    )


_ADDR_FIELDS = {
    1: {
        "street": (SystemField.ADDR1_STREET, COL_ADDR1_STREET),
        "house": (SystemField.ADDR1_HOUSE, COL_ADDR1_HOUSE),
        "apartment": (SystemField.ADDR1_APARTMENT, COL_ADDR1_APARTMENT),
        "entrance": (SystemField.ADDR1_ENTRANCE, COL_ADDR1_ENTRANCE),
        "floor": (SystemField.ADDR1_FLOOR, COL_ADDR1_FLOOR),
        "intercom": (SystemField.ADDR1_INTERCOM, COL_ADDR1_INTERCOM),
        "district_name": (SystemField.ADDR1_DISTRICT, COL_ADDR1_DISTRICT),
        "comment": (SystemField.ADDR1_COMMENT, COL_ADDR1_COMMENT),
    },
    2: {
        "street": (SystemField.ADDR2_STREET, COL_ADDR2_STREET),
        "house": (SystemField.ADDR2_HOUSE, COL_ADDR2_HOUSE),
        "apartment": (SystemField.ADDR2_APARTMENT, COL_ADDR2_APARTMENT),
        "entrance": (SystemField.ADDR2_ENTRANCE, COL_ADDR2_ENTRANCE),
        "floor": (SystemField.ADDR2_FLOOR, COL_ADDR2_FLOOR),
        "intercom": (SystemField.ADDR2_INTERCOM, COL_ADDR2_INTERCOM),
        "district_name": (SystemField.ADDR2_DISTRICT, COL_ADDR2_DISTRICT),
        "comment": (SystemField.ADDR2_COMMENT, COL_ADDR2_COMMENT),
    },
}


def _build_address(
    cells: dict[int, object],
    field_to_col: dict[SystemField, int] | None,
    street: str,
    house: str,
    *,
    addr_prefix: int,
) -> ParsedAddress:
    fields = _ADDR_FIELDS[addr_prefix]
    return ParsedAddress(
        street=street,
        house=house,
        apartment=_get(cells, field_to_col, *fields["apartment"]),
        entrance=_get(cells, field_to_col, *fields["entrance"]),
        floor=_get(cells, field_to_col, *fields["floor"]),
        intercom=_get(cells, field_to_col, *fields["intercom"]),
        district_name=_get(cells, field_to_col, *fields["district_name"]),
        comment=_get(cells, field_to_col, *fields["comment"]),
    )


def _build_optional_address(
    cells: dict[int, object],
    field_to_col: dict[SystemField, int] | None,
    *,
    addr_prefix: int,
) -> ParsedAddress | None:
    fields = _ADDR_FIELDS[addr_prefix]
    street = _get(cells, field_to_col, *fields["street"])
    house = _get(cells, field_to_col, *fields["house"])
    if not street or not house:
        return None
    return _build_address(
        cells, field_to_col, street, house, addr_prefix=addr_prefix
    )


def _get(
    cells: dict[int, object],
    field_to_col: dict[SystemField, int] | None,
    field: SystemField,
    legacy_col: int,
) -> str | None:
    if field_to_col is not None:
        col = field_to_col.get(field)
        if col is None:
            return None
        return _str(cells.get(col))
    return _str(cells.get(legacy_col))


def _str(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
