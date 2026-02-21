from io import BytesIO

from openpyxl import load_workbook

from backend.infrastructure.xlsx.client_error_xlsx_generator import (
    HEADERS,
    ClientErrorXlsxGenerator,
)
from backend.infrastructure.xlsx.client_xlsx_parser import (
    ImportErrorKind,
    RejectedClientRow,
)


def _make_rejected(
    row_number: int = 4,
    error_kind: ImportErrorKind = ImportErrorKind.EMPTY_FIELD,
    error_detail: str | None = None,
    raw_values: tuple[str | None, ...] | None = None,
) -> RejectedClientRow:
    if raw_values is None:
        raw_values = ("Іван", "+380501234567", None) + (None,) * 16
    return RejectedClientRow(
        row_number=row_number,
        raw_values=raw_values,
        error_kind=error_kind,
        error_detail=error_detail,
    )


def _parse_generated(xlsx_bytes: bytes):
    return load_workbook(BytesIO(xlsx_bytes), read_only=True).active


class TestClientErrorXlsxGenerator:
    def test_generate_with_rejected_rows(self):
        rows = [_make_rejected(row_number=4), _make_rejected(row_number=5)]
        xlsx_bytes = ClientErrorXlsxGenerator().generate(rows)

        ws = _parse_generated(xlsx_bytes)
        all_rows = list(ws.iter_rows(values_only=True))
        assert len(all_rows) == 3
        assert len(all_rows[0]) == 20

    def test_headers(self):
        xlsx_bytes = ClientErrorXlsxGenerator().generate([_make_rejected()])

        ws = _parse_generated(xlsx_bytes)
        header_row = [
            cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))
        ]
        assert len(header_row) == 20
        assert tuple(header_row) == HEADERS
        assert header_row[-1] == "Помилка"

    def test_raw_values_in_cells(self):
        raw = ("Іван", "+380501234567", None, "Хрещатик", "10") + (None,) * 14
        rows = [_make_rejected(raw_values=raw)]
        xlsx_bytes = ClientErrorXlsxGenerator().generate(rows)

        ws = _parse_generated(xlsx_bytes)
        data_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        assert data_row[0] == "Іван"
        assert data_row[1] == "+380501234567"
        assert data_row[3] == "Хрещатик"
        assert data_row[4] == "10"

    def test_error_message_in_last_column(self):
        rows = [_make_rejected(error_kind=ImportErrorKind.EMPTY_FIELD)]
        xlsx_bytes = ClientErrorXlsxGenerator().generate(rows)

        ws = _parse_generated(xlsx_bytes)
        data_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        assert data_row[19] == "Пусте поле"

    def test_empty_list(self):
        xlsx_bytes = ClientErrorXlsxGenerator().generate([])

        ws = _parse_generated(xlsx_bytes)
        all_rows = list(ws.iter_rows(values_only=True))
        assert len(all_rows) == 1

    def test_duplicate_in_file_error_message(self):
        rows = [
            _make_rejected(
                error_kind=ImportErrorKind.DUPLICATE_IN_FILE,
                error_detail="рядок 5",
            )
        ]
        xlsx_bytes = ClientErrorXlsxGenerator().generate(rows)

        ws = _parse_generated(xlsx_bytes)
        data_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        assert data_row[19] == "Дубль у файлі (рядок 5)"

    def test_bold_headers(self):
        xlsx_bytes = ClientErrorXlsxGenerator().generate([])

        wb = load_workbook(BytesIO(xlsx_bytes))
        ws = wb.active
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            assert cell.font.bold is True
