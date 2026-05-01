import uuid
from collections.abc import Callable
from datetime import UTC, datetime, time, timedelta
from decimal import Decimal
from typing import Any
from unittest.mock import AsyncMock, patch

import pytest
from fastapi import status
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.application.dto.coordinates import CoordinatesDTO
from backend.application.vars import ClientId, ShopRole
from backend.infrastructure.persistence.gateways.client_gateway import (
    SQLAlchemyClientGateway,
)
from backend.infrastructure.persistence.tables.clients import (
    Client,
    ClientAddress,
    ClientPhone,
)
from backend.infrastructure.persistence.tables.districts import District
from backend.infrastructure.persistence.tables.orders import Order

from .conftest import XLSX_CONTENT_TYPE

BASE_URL = "/api/v1/clients"


@pytest.mark.asyncio()
async def test_create_client_with_apartment(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1000
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Іван Іванов",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
                "apartment": "5",
                "entrance": "1",
                "floor": "2",
                "intercom": "5",
            }
        ],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()
    assert client_id is not None

    await session.flush()

    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["balance"] == 0

    await session.flush()

    result = await session.execute(
        select(Client).where(Client.id == uuid.UUID(client_id))
    )
    client = result.scalar_one()
    assert client.full_name == "Іван Іванов"

    phone_result = await session.execute(
        select(ClientPhone).where(
            ClientPhone.client_id == uuid.UUID(client_id)
        )
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380501234567"
    assert phones[0].is_primary is True

    address_result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].street == "Хрещатик"
    assert addresses[0].house == "10"
    assert addresses[0].apartment == "5"
    assert addresses[0].comment is None
    assert addresses[0].is_primary is True


@pytest.mark.asyncio()
async def test_create_client_with_coordinates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1010
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Клієнт з Координатами",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
                "apartment": "5",
                "coordinates": {
                    "latitude": 50.4501,
                    "longitude": 30.5234,
                },
            }
        ],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].latitude == 50.4501
    assert addresses[0].longitude == 30.5234


@pytest.mark.asyncio()
async def test_create_client_without_coordinates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1011
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Клієнт без Координат",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
            }
        ],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].latitude is None
    assert addresses[0].longitude is None


@pytest.mark.asyncio()
async def test_create_client_with_private_house(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1001
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Петро Петренко",
        "phones": [{"number": "+380931234567"}],
        "addresses": [
            {
                "street": "Заміська",
                "house": "15А",
                "comment": "Private house",
            }
        ],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()
    assert client_id is not None

    await session.flush()

    result = await session.execute(
        select(Client).where(Client.id == uuid.UUID(client_id))
    )
    client = result.scalar_one()
    assert client.full_name == "Петро Петренко"

    address_result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].street == "Заміська"
    assert addresses[0].house == "15А"
    assert addresses[0].comment == "Private house"
    assert addresses[0].apartment is None
    assert addresses[0].is_primary is True


@pytest.mark.asyncio()
async def test_create_client_with_multiple_phones_and_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1002
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Марія Марченко",
        "phones": [
            {"number": "+380671234567"},
            {"number": "+380631234567"},
            {"number": "+380501111111"},
        ],
        "addresses": [
            {
                "street": "Грушевського",
                "house": "5",
                "apartment": "10",
            },
            {
                "street": "Садова",
                "house": "22Б",
                "comment": "Private house",
            },
        ],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone)
        .where(ClientPhone.client_id == uuid.UUID(client_id))
        .order_by(ClientPhone.id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 3
    assert phones[0].number == "+380671234567"
    assert phones[0].is_primary is True
    assert phones[1].number == "+380631234567"
    assert phones[1].is_primary is False
    assert phones[2].number == "+380501111111"
    assert phones[2].is_primary is False

    address_result = await session.execute(
        select(ClientAddress)
        .where(ClientAddress.client_id == uuid.UUID(client_id))
        .order_by(ClientAddress.id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 2
    assert addresses[0].street == "Грушевського"
    assert addresses[0].is_primary is True
    assert addresses[1].street == "Садова"
    assert addresses[1].is_primary is False


@pytest.mark.asyncio()
async def test_create_client_without_phones_and_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1003
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Олексій Олексієнко",
        "phones": [],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    result = await session.execute(
        select(Client).where(Client.id == uuid.UUID(client_id))
    )
    client = result.scalar_one()
    assert client.full_name == "Олексій Олексієнко"

    phone_result = await session.execute(
        select(ClientPhone).where(
            ClientPhone.client_id == uuid.UUID(client_id)
        )
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 0

    address_result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 0


@pytest.mark.asyncio()
async def test_create_client_with_preferred_time_slot(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_time_slot,
) -> None:
    telegram_id = 1005
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)
    time_slot_id = await setup_test_time_slot(
        shop_id=shop_id,
        start_time=time(6, 0),
        end_time=time(8, 0),
        label="Ранній слот",
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=BASE_URL,
        headers=headers,
        json={
            "full_name": "Клієнт з улюбленим слотом",
            "phones": [{"number": "+380501234567"}],
            "addresses": [],
            "preferred_time_slot_id": str(time_slot_id),
        },
    )

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["preferred_time_slot_id"] == str(
        time_slot_id
    )


@pytest.mark.asyncio()
async def test_create_client_unauthorized(
    http_client: AsyncClient,
) -> None:
    json = {
        "full_name": "Test User",
        "phones": [{"number": "+380501234567"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, json=json)

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_create_client_as_courier_forbidden(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1004
    await setup_full_test_user_with_shop(
        telegram_id=telegram_id, role=ShopRole.COURIER
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Test User",
        "phones": [{"number": "+380501234567"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio()
async def test_get_client_by_id(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2000
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create client directly in DB
    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Тестовий Клієнт",
        phones=["+380501234567", "+380507654321"],
        addresses=[
            {
                "street": "Хрещатик",
                "house": "10",
                "apartment": "5",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    # Get the client by ID
    response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data["client_id"] == str(client_id)
    assert data["full_name"] == "Тестовий Клієнт"
    assert len(data["phones"]) == 2
    assert data["phones"][0]["number"] == "+380501234567"
    assert data["phones"][0]["is_primary"] is True
    assert data["phones"][1]["number"] == "+380507654321"
    assert data["phones"][1]["is_primary"] is False
    assert len(data["addresses"]) == 1
    assert data["addresses"][0]["street"] == "Хрещатик"
    assert data["addresses"][0]["is_primary"] is True


@pytest.mark.asyncio()
async def test_get_client_returns_balance(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2002
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Балансовий клієнт",
        balance=Decimal("1250.75"),
        phones=["+380501234567"],
        addresses=[{"street": "Хрещатик", "house": "10"}],
    )
    await session.commit()

    response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=customer_headers(telegram_id)
    )

    assert response.status_code == status.HTTP_200_OK
    balance = response.json()["balance"]
    assert isinstance(balance, float)
    assert balance == 1250.75


@pytest.mark.asyncio()
async def test_set_client_balance(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2004
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт для балансу",
        balance=Decimal("100.00"),
    )
    await session.commit()

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}/balance",
        headers=customer_headers(telegram_id),
        json={"balance": 250.75},
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()
    stored_client = await session.get(Client, client_id)
    assert stored_client is not None
    assert stored_client.balance == Decimal("250.75")

    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}",
        headers=customer_headers(telegram_id),
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["balance"] == 250.75


@pytest.mark.asyncio()
async def test_set_client_balance_rejects_over_precision(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2006
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(shop_id=shop_id)
    await session.commit()

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}/balance",
        headers=customer_headers(telegram_id),
        json={"balance": 1.999},
    )

    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT


@pytest.mark.asyncio()
async def test_set_client_balance_as_courier_forbidden(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2007
    _, shop_id = await setup_full_test_user_with_shop(
        telegram_id=telegram_id, role=ShopRole.COURIER
    )

    client_id = await setup_test_client(shop_id=shop_id)
    await session.commit()

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}/balance",
        headers=customer_headers(telegram_id),
        json={"balance": 250.75},
    )

    assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio()
async def test_set_client_balance_not_found(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 2005
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    response = await http_client.patch(
        url=f"{BASE_URL}/{uuid.uuid4()}/balance",
        headers=customer_headers(telegram_id),
        json={"balance": 250.75},
    )

    assert response.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.asyncio()
async def test_bulk_save_client_defaults_balance_to_zero(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 2003
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = ClientId(uuid.uuid4())
    client = Client(id=client_id, shop_id=shop_id, full_name="Bulk Client")
    client.balance = None
    gateway = SQLAlchemyClientGateway(session)

    await gateway.save_all([client])
    await session.commit()

    response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=customer_headers(telegram_id)
    )

    assert response.status_code == status.HTTP_200_OK
    balance = response.json()["balance"]
    assert isinstance(balance, (int, float))
    assert balance == 0


@pytest.mark.asyncio()
async def test_get_client_by_id_not_found(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 2001
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    non_existent_id = str(uuid.uuid4())

    response = await http_client.get(
        url=f"{BASE_URL}/{non_existent_id}", headers=headers
    )

    assert response.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.asyncio()
async def test_get_client_unauthorized(
    http_client: AsyncClient,
) -> None:
    client_id = str(uuid.uuid4())

    response = await http_client.get(url=f"{BASE_URL}/{client_id}")

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_get_all_clients(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2100
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create three clients directly in DB
    await setup_test_client(
        shop_id=shop_id,
        full_name="Анна Антоненко",
        balance=Decimal("14.50"),
        phones=["+380501111111"],
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Борис Борисенко",
        phones=["+380502222222"],
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Віктор Вікторенко",
        phones=["+380503333333"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    # Get all clients
    response = await http_client.get(url=f"{BASE_URL}/all", headers=headers)

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 3
    assert isinstance(data[0]["balance"], float)
    assert data[0]["balance"] == 14.5
    assert data[0]["full_name"] == "Анна Антоненко"
    assert data[1]["full_name"] == "Борис Борисенко"
    assert data[2]["full_name"] == "Віктор Вікторенко"


@pytest.mark.asyncio()
async def test_get_all_clients_returns_preferred_time_slot(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
    setup_test_time_slot,
) -> None:
    telegram_id = 2104
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)
    time_slot_id = await setup_test_time_slot(
        shop_id=shop_id,
        start_time=time(6, 0),
        end_time=time(8, 0),
        label="Ранній слот",
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт зі слотом",
        preferred_time_slot_id=time_slot_id,
    )
    await session.commit()

    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=customer_headers(telegram_id)
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 1
    assert data[0]["preferred_time_slot_id"] == str(time_slot_id)


@pytest.mark.asyncio()
async def test_get_all_clients_filter_by_full_name(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2101
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create clients directly in DB
    await setup_test_client(
        shop_id=shop_id,
        full_name="Анна Антоненко",
        phones=["+380501111111"],
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Борис Борисенко",
        phones=["+380502222222"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    # Filter by full name
    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=headers, params={"full_name": "Анна"}
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 1
    assert data[0]["full_name"] == "Анна Антоненко"


@pytest.mark.asyncio()
async def test_get_all_clients_filter_by_phone(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2102
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create clients directly in DB
    await setup_test_client(
        shop_id=shop_id,
        full_name="Анна Антоненко",
        phones=["+380501111111"],
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Борис Борисенко",
        phones=["+380502222222"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    # Filter by phone
    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=headers, params={"phone": "1111"}
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 1
    assert data[0]["full_name"] == "Анна Антоненко"


@pytest.mark.asyncio()
async def test_get_all_clients_with_pagination(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2104
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create 5 clients directly in DB
    for i in range(5):
        await setup_test_client(
            shop_id=shop_id,
            full_name=f"Client {i:02d}",
        )
    await session.commit()

    headers = customer_headers(telegram_id)

    # Get first 3 clients
    response = await http_client.get(
        url=f"{BASE_URL}/all",
        headers=headers,
        params={"limit": 3, "offset": 0},
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 3

    # Get next 2 clients
    response = await http_client.get(
        url=f"{BASE_URL}/all",
        headers=headers,
        params={"limit": 3, "offset": 3},
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data) == 2


@pytest.mark.asyncio()
async def test_get_all_clients_with_sorting(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2105
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    # Create clients directly in DB
    await setup_test_client(shop_id=shop_id, full_name="Віктор")
    await setup_test_client(shop_id=shop_id, full_name="Анна")
    await setup_test_client(shop_id=shop_id, full_name="Борис")
    await session.commit()

    headers = customer_headers(telegram_id)

    # Sort ascending (default)
    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=headers, params={"order": "ASC"}
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data[0]["full_name"] == "Анна"
    assert data[1]["full_name"] == "Борис"
    assert data[2]["full_name"] == "Віктор"

    # Sort descending
    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=headers, params={"order": "DESC"}
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data[0]["full_name"] == "Віктор"
    assert data[1]["full_name"] == "Борис"
    assert data[2]["full_name"] == "Анна"


@pytest.mark.asyncio()
async def test_get_all_clients_unauthorized(
    http_client: AsyncClient,
) -> None:
    response = await http_client.get(url=f"{BASE_URL}/all")

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_delete_product(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 1000
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)
    client_id = await setup_test_client(shop_id=shop_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    url = BASE_URL + f"/{client_id}"
    response = await http_client.delete(url=url, headers=headers)

    assert response.status_code == status.HTTP_204_NO_CONTENT
    await session.flush()

    deleted_entity = await session.execute(
        select(Client).where(Client.id == client_id)
    )
    assert deleted_entity.scalar_one_or_none() is None


@pytest.mark.asyncio()
async def test_edit_client_full_name(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3000
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Оригінальне Ім'я",
        phones=["+380501111111", "+380502222222"],
        addresses=[
            {
                "street": "Оригінальна вулиця",
                "house": "1",
                "apartment": "10",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {"full_name": "Оновлене Ім'я"}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    result = await session.execute(
        select(Client).where(Client.id == client_id)
    )
    client = result.scalar_one()
    assert client.full_name == "Оновлене Ім'я"


@pytest.mark.asyncio()
async def test_edit_client_balance(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3001
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Балансовий клієнт",
        balance=Decimal("100.00"),
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
        json={"balance": -250.75},
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    stored_client = await session.get(Client, client_id)
    assert stored_client is not None
    assert stored_client.balance == Decimal("-250.75")

    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["balance"] == -250.75


@pytest.mark.asyncio()
async def test_edit_client_preferred_time_slot_set_change_and_clear(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
    setup_test_time_slot,
) -> None:
    telegram_id = 3012
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)
    morning_slot_id = await setup_test_time_slot(
        shop_id=shop_id,
        start_time=time(6, 0),
        end_time=time(8, 0),
        label="Ранній слот",
    )
    evening_slot_id = await setup_test_time_slot(
        shop_id=shop_id,
        start_time=time(21, 0),
        end_time=time(23, 0),
        label="Пізній слот",
    )
    client_id = await setup_test_client(shop_id=shop_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
        json={"preferred_time_slot_id": str(morning_slot_id)},
    )
    assert response.status_code == status.HTTP_200_OK
    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["preferred_time_slot_id"] == str(
        morning_slot_id
    )

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
        json={"preferred_time_slot_id": str(evening_slot_id)},
    )
    assert response.status_code == status.HTTP_200_OK
    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["preferred_time_slot_id"] == str(
        evening_slot_id
    )

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
        json={"preferred_time_slot_id": None},
    )
    assert response.status_code == status.HTTP_200_OK
    client_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert client_response.status_code == status.HTTP_200_OK
    assert client_response.json()["preferred_time_slot_id"] is None


@pytest.mark.asyncio()
async def test_edit_client_rejects_preferred_time_slot_from_another_shop(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
    setup_test_time_slot,
    create_user,
    create_telegram_account,
    create_shop,
    create_shop_membership,
) -> None:
    telegram_id = 3013
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)
    client_id = await setup_test_client(shop_id=shop_id)

    other_user_id = await create_user()
    await create_telegram_account(
        user_id=other_user_id, telegram_id=3014, full_name="Other User"
    )
    other_shop_id = await create_shop()
    await create_shop_membership(user_id=other_user_id, shop_id=other_shop_id)
    other_time_slot_id = await setup_test_time_slot(
        shop_id=other_shop_id,
        start_time=time(6, 0),
        end_time=time(8, 0),
        label="Чужий слот",
    )
    await session.commit()

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=customer_headers(telegram_id),
        json={"preferred_time_slot_id": str(other_time_slot_id)},
    )

    assert response.status_code == status.HTTP_403_FORBIDDEN
    stored_client = await session.get(Client, client_id)
    assert stored_client is not None
    assert stored_client.preferred_time_slot_id is None


@pytest.mark.asyncio()
async def test_edit_client_phones(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3002
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Ім'я Клієнта",
        phones=["+380501111111", "+380502222222"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "phones": [
            {"number": "+380509999999", "is_primary": True},
            {"number": "+380508888888", "is_primary": False},
            {"number": "+380507777777", "is_primary": False},
        ]
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )
    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone)
        .where(ClientPhone.client_id == client_id)
        .order_by(ClientPhone.id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 3
    assert phones[0].number == "+380509999999"
    assert phones[0].is_primary is True
    assert phones[1].number == "+380508888888"
    assert phones[1].is_primary is False
    assert phones[2].number == "+380507777777"
    assert phones[2].is_primary is False


@pytest.mark.asyncio()
async def test_edit_client_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3003
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Ім'я Клієнта",
        addresses=[
            {
                "street": "Стара вулиця",
                "house": "1",
                "apartment": "10",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "addresses": [
            {
                "street": "Нова вулиця",
                "house": "100",
                "comment": "Private house",
                "is_primary": True,
            },
            {
                "street": "Ще одна вулиця",
                "house": "200",
                "apartment": "50",
                "is_primary": False,
            },
        ]
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress)
        .where(ClientAddress.client_id == client_id)
        .order_by(ClientAddress.id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 2
    assert addresses[0].street == "Нова вулиця"
    assert addresses[0].house == "100"
    assert addresses[0].comment == "Private house"
    assert addresses[0].is_primary is True
    assert addresses[1].street == "Ще одна вулиця"
    assert addresses[1].is_primary is False


@pytest.mark.asyncio()
async def test_edit_client_add_coordinates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3009
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        addresses=[
            {
                "street": "Вулиця",
                "house": "1",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "addresses": [
            {
                "street": "Вулиця",
                "house": "1",
                "coordinates": {
                    "latitude": 50.4501,
                    "longitude": 30.5234,
                },
                "is_primary": True,
            }
        ]
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].latitude == 50.4501
    assert addresses[0].longitude == 30.5234


@pytest.mark.asyncio()
async def test_edit_client_remove_coordinates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3010
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        addresses=[
            {
                "street": "Вулиця",
                "house": "1",
                "latitude": 50.4501,
                "longitude": 30.5234,
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "addresses": [
            {
                "street": "Вулиця",
                "house": "1",
                "is_primary": True,
            }
        ]
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].latitude is None
    assert addresses[0].longitude is None


@pytest.mark.asyncio()
async def test_get_client_coordinates_in_response(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3011
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт з Координатами",
        phones=["+380501234567"],
        addresses=[
            {
                "street": "Хрещатик",
                "house": "10",
                "latitude": 50.4501,
                "longitude": 30.5234,
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert len(data["addresses"]) == 1
    assert data["addresses"][0]["coordinates"]["latitude"] == 50.4501
    assert data["addresses"][0]["coordinates"]["longitude"] == 30.5234


@pytest.mark.asyncio()
async def test_edit_client_all_fields(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3004
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Старе Ім'я",
        phones=["+380501111111"],
        addresses=[
            {
                "street": "Стара вулиця",
                "house": "1",
                "apartment": "10",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Повністю Нове Ім'я",
        "phones": [{"number": "+380501234567", "is_primary": True}],
        "addresses": [
            {
                "street": "Повністю нова вулиця",
                "house": "999",
                "comment": "Private house",
                "is_primary": True,
            }
        ],
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    result = await session.execute(
        select(Client).where(Client.id == client_id)
    )
    client = result.scalar_one()
    assert client.full_name == "Повністю Нове Ім'я"

    phone_result = await session.execute(
        select(ClientPhone).where(ClientPhone.client_id == client_id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380501234567"

    address_result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].street == "Повністю нова вулиця"


@pytest.mark.asyncio()
async def test_edit_client_clear_phones(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3005
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Ім'я Клієнта",
        phones=["+380501111111", "+380502222222"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json: dict[str, list[Any]] = {"phones": []}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone).where(ClientPhone.client_id == client_id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 0


@pytest.mark.asyncio()
async def test_edit_client_clear_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3006
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Ім'я Клієнта",
        addresses=[
            {
                "street": "Вулиця",
                "house": "1",
                "apartment": "10",
            }
        ],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json: dict[str, list[Any]] = {"addresses": []}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    address_result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    addresses = address_result.scalars().all()
    assert len(addresses) == 0


@pytest.mark.asyncio()
async def test_edit_client_not_found(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 3007
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    non_existent_id = str(uuid.uuid4())

    json = {"full_name": "Нове Ім'я"}

    response = await http_client.patch(
        url=f"{BASE_URL}/{non_existent_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.asyncio()
async def test_edit_client_unauthorized(
    http_client: AsyncClient,
) -> None:
    client_id = str(uuid.uuid4())

    json = {"full_name": "Нове Ім'я"}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", json=json
    )

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_edit_client_as_courier_forbidden(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3008
    _, shop_id = await setup_full_test_user_with_shop(
        telegram_id=telegram_id, role=ShopRole.COURIER
    )

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Ім'я Клієнта",
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {"full_name": "Нове Ім'я"}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio()
async def test_create_client_with_duplicate_phone_returns_409(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 3100
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json1 = {
        "full_name": "Перший Клієнт",
        "phones": [{"number": "+380991234567"}],
        "addresses": [],
    }

    response1 = await http_client.post(
        url=BASE_URL, headers=headers, json=json1
    )
    assert response1.status_code == status.HTTP_201_CREATED

    await session.flush()

    json2 = {
        "full_name": "Другий Клієнт",
        "phones": [{"number": "+380991234567"}],
        "addresses": [],
    }

    response2 = await http_client.post(
        url=BASE_URL, headers=headers, json=json2
    )

    assert response2.status_code == status.HTTP_409_CONFLICT
    data = response2.json()
    assert data["code"] == "duplicate_phones"
    assert len(data["duplicates"]) == 1
    assert data["duplicates"][0]["phone_number"] == "+380991234567"
    assert (
        data["duplicates"][0]["existing_clients"][0]["full_name"]
        == "Перший Клієнт"
    )


@pytest.mark.asyncio()
async def test_create_client_with_duplicate_phone_confirmed(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 3103
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json1 = {
        "full_name": "Перший Клієнт",
        "phones": [{"number": "+380991234567"}],
        "addresses": [],
    }

    response1 = await http_client.post(
        url=BASE_URL, headers=headers, json=json1
    )
    assert response1.status_code == status.HTTP_201_CREATED

    await session.flush()

    json2 = {
        "full_name": "Другий Клієнт",
        "phones": [{"number": "+380991234567"}],
        "addresses": [],
        "confirm_duplicate_phones": True,
    }

    response2 = await http_client.post(
        url=BASE_URL, headers=headers, json=json2
    )

    assert response2.status_code == status.HTTP_201_CREATED
    assert response2.json() is not None


@pytest.mark.asyncio()
async def test_edit_client_with_duplicate_phone_returns_409(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3101
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    await setup_test_client(
        shop_id=shop_id,
        full_name="Перший Клієнт",
        phones=["+380509999999"],
    )

    client2_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Другий Клієнт",
        phones=["+380508888888"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {"phones": [{"number": "+380509999999", "is_primary": True}]}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client2_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_409_CONFLICT
    data = response.json()
    assert data["code"] == "duplicate_phones"
    assert data["duplicates"][0]["phone_number"] == "+380509999999"
    assert (
        data["duplicates"][0]["existing_clients"][0]["full_name"]
        == "Перший Клієнт"
    )


@pytest.mark.asyncio()
async def test_edit_client_with_duplicate_phone_confirmed(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3104
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    await setup_test_client(
        shop_id=shop_id,
        full_name="Перший Клієнт",
        phones=["+380509999999"],
    )

    client2_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Другий Клієнт",
        phones=["+380508888888"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "phones": [{"number": "+380509999999", "is_primary": True}],
        "confirm_duplicate_phones": True,
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client2_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK


@pytest.mark.asyncio()
async def test_edit_client_keep_same_phone_numbers_integration(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 3102
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        phones=["+380501111111", "+380502222222"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    # First, get the client to obtain phone IDs
    get_response = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    assert get_response.status_code == status.HTTP_200_OK
    client_data = get_response.json()
    phone_ids = [phone["id"] for phone in client_data["phones"]]

    json = {
        "phones": [
            {
                "number": "+380501111111",
                "is_primary": True,
                "id": phone_ids[0],
            },
            {
                "number": "+380502222222",
                "is_primary": False,
                "id": phone_ids[1],
            },
        ]
    }

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone)
        .where(ClientPhone.client_id == client_id)
        .order_by(ClientPhone.id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 2
    assert phones[0].number == "+380501111111"
    assert phones[0].is_primary is True
    assert phones[1].number == "+380502222222"
    assert phones[1].is_primary is False


@pytest.mark.asyncio()
async def test_get_all_clients_filter_by_phone_no_duplicates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2150
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт з багатьма телефонами",
        phones=[
            "+380501234567",
            "+380501234568",
            "+380501234569",
        ],
    )
    await setup_test_client(
        shop_id=shop_id,
        full_name="Інший Клієнт",
        phones=["+380509999999"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.get(
        url=f"{BASE_URL}/all", headers=headers, params={"phone": "50123456"}
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()

    assert len(data) == 1
    assert data[0]["full_name"] == "Клієнт з багатьма телефонами"


@pytest.mark.asyncio()
async def test_get_all_clients_pagination_no_duplicates_same_name(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 2160
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    for _i in range(10):
        await setup_test_client(
            shop_id=shop_id,
            full_name="Однакове Ім'я",
        )
    await session.commit()

    headers = customer_headers(telegram_id)

    all_client_ids: list[str] = []
    for offset in range(0, 10, 2):
        response = await http_client.get(
            url=f"{BASE_URL}/all",
            headers=headers,
            params={"limit": 2, "offset": offset},
        )
        assert response.status_code == status.HTTP_200_OK
        page_ids = [c["client_id"] for c in response.json()]
        all_client_ids.extend(page_ids)

    assert len(all_client_ids) == 10
    assert len(all_client_ids) == len(set(all_client_ids))


@pytest.mark.asyncio()
async def test_create_client_with_phone_normalization_leading_zero(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 4000
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Тест Нормалізації",
        "phones": [{"number": "0980074978"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone).where(
            ClientPhone.client_id == uuid.UUID(client_id)
        )
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380980074978"


@pytest.mark.asyncio()
async def test_create_client_with_phone_normalization_with_spaces(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 4001
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Тест Нормалізації з Пробілами",
        "phones": [{"number": "098 007 49 78"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone).where(
            ClientPhone.client_id == uuid.UUID(client_id)
        )
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380980074978"


@pytest.mark.asyncio()
async def test_create_client_with_phone_normalization_with_formatting(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 4002
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Тест Нормалізації з Форматуванням",
        "phones": [{"number": "+38(098)007-49-78"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_201_CREATED
    client_id = response.json()

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone).where(
            ClientPhone.client_id == uuid.UUID(client_id)
        )
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380980074978"


@pytest.mark.asyncio()
async def test_create_client_with_invalid_phone_too_short(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 4003
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Тест Невалідного Номеру",
        "phones": [{"number": "098007497"}],
        "addresses": [],
    }

    response = await http_client.post(url=BASE_URL, headers=headers, json=json)

    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert "phone" in response.json()["detail"].lower()


@pytest.mark.asyncio()
async def test_edit_client_with_phone_normalization(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 4004
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        phones=["+380501111111"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {"phones": [{"number": "0502222222", "is_primary": True}]}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    phone_result = await session.execute(
        select(ClientPhone).where(ClientPhone.client_id == client_id)
    )
    phones = phone_result.scalars().all()
    assert len(phones) == 1
    assert phones[0].number == "+380502222222"


@pytest.mark.asyncio()
async def test_edit_client_with_invalid_phone(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 4005
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        phones=["+380501111111"],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {"phones": [{"number": "123", "is_primary": True}]}

    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}", headers=headers, json=json
    )

    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert "phone" in response.json()["detail"].lower()


@pytest.mark.asyncio()
async def test_delete_client_unauthorized(
    http_client: AsyncClient,
) -> None:
    client_id = str(uuid.uuid4())
    response = await http_client.delete(url=f"{BASE_URL}/{client_id}")

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_delete_client_as_courier_forbidden(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 1050
    _, shop_id = await setup_full_test_user_with_shop(
        telegram_id=telegram_id, role=ShopRole.COURIER
    )
    client_id = await setup_test_client(shop_id=shop_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    url = BASE_URL + f"/{client_id}"
    response = await http_client.delete(url=url, headers=headers)

    assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio()
async def test_delete_client_not_found_returns_ok(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1051
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    url = BASE_URL + f"/{uuid.uuid4()}"
    response = await http_client.delete(url=url, headers=headers)

    assert response.status_code == status.HTTP_204_NO_CONTENT


GEOCODER_GEOCODE = "backend.application.services.geocoder.Geocoder.geocode"


@pytest.mark.asyncio()
async def test_create_client_geocodes_address_when_shop_has_city(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1060
    await setup_full_test_user_with_shop(
        telegram_id=telegram_id, shop_city="Київ"
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Геокод Клієнт",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
            }
        ],
    }

    fake_coords = CoordinatesDTO(latitude=50.45, longitude=30.52)

    with patch(
        GEOCODER_GEOCODE,
        new_callable=AsyncMock,
        return_value=fake_coords,
    ) as mock_geocode:
        response = await http_client.post(
            url=BASE_URL, headers=headers, json=json
        )

        assert response.status_code == status.HTTP_201_CREATED
        mock_geocode.assert_called_once_with(
            "Хрещатик", "10", "Київ", require_house=True
        )

    client_id = response.json()
    await session.flush()

    result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addr = result.scalar_one()
    assert addr.latitude == 50.45
    assert addr.longitude == 30.52


@pytest.mark.asyncio()
async def test_create_client_skips_geocoding_when_coords_provided(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1061
    await setup_full_test_user_with_shop(
        telegram_id=telegram_id, shop_city="Київ"
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Клієнт з Координатами",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
                "coordinates": {
                    "latitude": 50.4501,
                    "longitude": 30.5234,
                },
            }
        ],
    }

    with patch(
        GEOCODER_GEOCODE,
        new_callable=AsyncMock,
    ) as mock_geocode:
        response = await http_client.post(
            url=BASE_URL, headers=headers, json=json
        )

        assert response.status_code == status.HTTP_201_CREATED
        mock_geocode.assert_not_called()

    client_id = response.json()
    await session.flush()

    result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addr = result.scalar_one()
    assert addr.latitude == 50.4501
    assert addr.longitude == 30.5234


@pytest.mark.asyncio()
async def test_create_client_skips_geocoding_when_no_city(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 1062
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    json = {
        "full_name": "Клієнт без Міста",
        "phones": [{"number": "+380501234567"}],
        "addresses": [
            {
                "street": "Хрещатик",
                "house": "10",
            }
        ],
    }

    with patch(
        GEOCODER_GEOCODE,
        new_callable=AsyncMock,
    ) as mock_geocode:
        response = await http_client.post(
            url=BASE_URL, headers=headers, json=json
        )

        assert response.status_code == status.HTTP_201_CREATED
        mock_geocode.assert_not_called()

    client_id = response.json()
    await session.flush()

    result = await session.execute(
        select(ClientAddress).where(
            ClientAddress.client_id == uuid.UUID(client_id)
        )
    )
    addr = result.scalar_one()
    assert addr.latitude is None
    assert addr.longitude is None


@pytest.mark.asyncio()
async def test_edit_client_geocodes_new_address(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
) -> None:
    telegram_id = 1063
    _, shop_id = await setup_full_test_user_with_shop(
        telegram_id=telegram_id, shop_city="Київ"
    )

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт",
        phones=["+380501234567"],
        addresses=[{"street": "Стара", "house": "1"}],
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    client_resp = await http_client.get(
        url=f"{BASE_URL}/{client_id}", headers=headers
    )
    client_data = client_resp.json()
    phone_id = client_data["phones"][0]["id"]

    fake_coords = CoordinatesDTO(latitude=50.46, longitude=30.53)

    json = {
        "phones": [
            {
                "number": "+380501234567",
                "is_primary": True,
                "id": phone_id,
            }
        ],
        "addresses": [
            {
                "street": "Нова вулиця",
                "house": "5",
                "is_primary": True,
            }
        ],
    }

    with patch(
        GEOCODER_GEOCODE,
        new_callable=AsyncMock,
        return_value=fake_coords,
    ) as mock_geocode:
        response = await http_client.patch(
            url=f"{BASE_URL}/{client_id}",
            headers=headers,
            json=json,
        )

        assert response.status_code == status.HTTP_200_OK
        mock_geocode.assert_called_once_with(
            "Нова вулиця", "5", "Київ", require_house=True
        )

    await session.flush()

    result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    addresses = result.scalars().all()
    assert len(addresses) == 1
    assert addresses[0].latitude == 50.46
    assert addresses[0].longitude == 30.53


IMPORT_URL = f"{BASE_URL}/import"


@pytest.mark.asyncio()
async def test_import_clients_success(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5000
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Іван Іванов", "+380501234567", None, "Хрещатик", "10", "5"],
            ["Петро Петренко", "0931234567", "+380671111111", "Садова", "22"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 2
    assert response.json()["error_file_id"] is None

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2

    names = sorted(c.full_name for c in clients)
    assert names == ["Іван Іванов", "Петро Петренко"]


@pytest.mark.asyncio()
async def test_import_clients_rejects_non_xlsx(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
) -> None:
    telegram_id = 5001
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files={"file": ("clients.csv", b"data", "text/csv")},
    )

    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT


@pytest.mark.asyncio()
async def test_import_clients_unauthorized(
    http_client: AsyncClient,
    build_xlsx: Callable[..., bytes],
) -> None:
    response = await http_client.post(
        url=IMPORT_URL,
        files={"file": ("clients.xlsx", build_xlsx([]), XLSX_CONTENT_TYPE)},
    )

    assert response.status_code == status.HTTP_401_UNAUTHORIZED


@pytest.mark.asyncio()
async def test_import_clients_as_courier_forbidden(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5002
    await setup_full_test_user_with_shop(
        telegram_id=telegram_id, role=ShopRole.COURIER
    )
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Тест", "+380501234567", None, "Вулиця", "1"],
        ]),
    )

    assert response.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.asyncio()
async def test_import_clients_empty_file(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5003
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 0
    assert response.json()["error_file_id"] is None


@pytest.mark.asyncio()
async def test_import_clients_skips_row_with_invalid_phone(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5004
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Валідний", "+380501234567", None, "Вулиця", "1"],
            ["Невалідний", "123", None, "Вулиця", "2"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 1
    assert response.json()["error_file_id"] is not None

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 1
    assert clients[0].full_name == "Валідний"


@pytest.mark.asyncio()
async def test_import_clients_with_two_phones_and_two_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5005
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [
                "Клієнт",
                "+380501234567",
                "+380671111111",
                "Хрещатик",
                "10",
                "5",
                "1",
                "3",
                "5",
                None,
                "Коментар 1",
                "Садова",
                "22",
                None,
                None,
                None,
                None,
                None,
                "Приватний будинок",
            ],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 1

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 1
    client = clients[0]

    phones = (
        (
            await session.execute(
                select(ClientPhone)
                .where(ClientPhone.client_id == client.id)
                .order_by(ClientPhone.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(phones) == 2
    assert phones[0].number == "+380501234567"
    assert phones[0].is_primary is True
    assert phones[1].number == "+380671111111"
    assert phones[1].is_primary is False

    addresses = (
        (
            await session.execute(
                select(ClientAddress)
                .where(ClientAddress.client_id == client.id)
                .order_by(ClientAddress.id)
            )
        )
        .scalars()
        .all()
    )
    assert len(addresses) == 2
    assert addresses[0].street == "Хрещатик"
    assert addresses[0].apartment == "5"
    assert addresses[0].comment == "Коментар 1"
    assert addresses[0].is_primary is True
    assert addresses[1].street == "Садова"
    assert addresses[1].comment == "Приватний будинок"
    assert addresses[1].is_primary is False


@pytest.mark.asyncio()
async def test_import_clients_creates_districts(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5006
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [
                "Клієнт 1",
                "+380501234567",
                None,
                "Хрещатик",
                "10",
                None,
                None,
                None,
                None,
                "Центр",
            ],
            [
                "Клієнт 2",
                "+380502222222",
                None,
                "Садова",
                "5",
                None,
                None,
                None,
                None,
                "Центр",
            ],
            [
                "Клієнт 3",
                "+380503333333",
                None,
                "Заміська",
                "1",
                None,
                None,
                None,
                None,
                "Околиця",
            ],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 3

    await session.flush()

    districts = (await session.execute(select(District))).scalars().all()
    district_names = sorted(d.name for d in districts)
    assert district_names == ["Околиця", "Центр"]


@pytest.mark.asyncio()
async def test_import_clients_skips_row_without_required_fields(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5007
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [None, "+380501234567", None, "Вулиця", "1"],
            ["Клієнт", None, None, "Вулиця", "1"],
            ["Клієнт", "+380501234567", None, None, "1"],
            ["Клієнт", "+380501234567", None, "Вулиця", None],
            ["Валідний", "+380509999999", None, "Вулиця", "1"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 1
    assert response.json()["error_file_id"] is not None


@pytest.mark.asyncio()
async def test_import_clients_skips_duplicates_on_reupload(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5010
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)
    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ["Петро Петренко", "+380931234567", None, "Садова", "22"],
    ])

    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 2
    assert response1.json()["skipped"] == 0

    await session.flush()

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ["Петро Петренко", "+380931234567", None, "Садова", "22"],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 0
    assert response2.json()["skipped"] == 2
    assert response2.json()["error_file_id"] is not None

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2


@pytest.mark.asyncio()
async def test_import_clients_skips_duplicates_with_swapped_phones(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5011
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        ["Клієнт", "+380501111111", "+380502222222", "Хрещатик", "10"],
    ])
    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 1

    await session.flush()

    files = upload_xlsx([
        ["Клієнт", "+380502222222", "+380501111111", "Хрещатик", "10"],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 0
    assert response2.json()["skipped"] == 1

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 1


@pytest.mark.asyncio()
async def test_import_clients_skips_duplicates_with_swapped_addresses(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5012
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        [
            "Клієнт",
            "+380501111111",
            None,
            "Хрещатик",
            "10",
            None,
            None,
            None,
            None,
            None,
            None,
            "Садова",
            "22",
        ],
    ])
    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 1

    await session.flush()

    files = upload_xlsx([
        [
            "Клієнт",
            "+380501111111",
            None,
            "Садова",
            "22",
            None,
            None,
            None,
            None,
            None,
            None,
            "Хрещатик",
            "10",
        ],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 0
    assert response2.json()["skipped"] == 1

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 1


@pytest.mark.asyncio()
async def test_import_clients_skips_intra_file_duplicates(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5013
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ["Петро Петренко", "+380931234567", None, "Садова", "22"],
    ])
    response = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 2
    assert response.json()["skipped"] == 1
    assert response.json()["error_file_id"] is not None

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2


@pytest.mark.asyncio()
async def test_import_clients_imports_different_client_same_phone(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5014
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
    ])
    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 1

    await session.flush()

    files = upload_xlsx([
        ["Петро Петренко", "+380501234567", None, "Хрещатик", "10"],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 1
    assert response2.json()["skipped"] == 0

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2


@pytest.mark.asyncio()
async def test_import_clients_imports_same_name_different_address(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5015
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
    ])
    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 1

    await session.flush()

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Садова", "5"],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 1
    assert response2.json()["skipped"] == 0

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2


@pytest.mark.asyncio()
async def test_import_clients_partial_reupload(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5016
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
    ])
    response1 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response1.status_code == status.HTTP_200_OK
    assert response1.json()["imported"] == 1

    await session.flush()

    files = upload_xlsx([
        ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ["Новий Клієнт", "+380939999999", None, "Нова", "1"],
    ])
    response2 = await http_client.post(
        url=IMPORT_URL, headers=headers, files=files
    )
    assert response2.status_code == status.HTTP_200_OK
    assert response2.json()["imported"] == 1
    assert response2.json()["skipped"] == 1

    await session.flush()

    clients = (await session.execute(select(Client))).scalars().all()
    assert len(clients) == 2


ERRORS_EXPORT_URL = f"{BASE_URL}/export/errors"


@pytest.mark.asyncio()
async def test_import_returns_error_file_on_invalid_rows(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5020
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Валідний", "+380501234567", None, "Хрещатик", "10"],
            [None, "+380502222222", None, "Садова", "5"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data["imported"] == 1
    assert data["error_file_id"] is not None
    assert data["error_filename"] == "import_errors.xlsx"


@pytest.mark.asyncio()
async def test_import_returns_no_error_file_on_success(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5021
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Клієнт 1", "+380501234567", None, "Хрещатик", "10"],
            ["Клієнт 2", "+380931234567", None, "Садова", "22"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["error_file_id"] is None


@pytest.mark.asyncio()
async def test_import_all_invalid_returns_error_file(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5022
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [None, "+380501234567", None, "Хрещатик", "10"],
            ["Клієнт", None, None, "Садова", "5"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    data = response.json()
    assert data["imported"] == 0
    assert data["skipped"] == 2
    assert data["error_file_id"] is not None


@pytest.mark.asyncio()
async def test_download_import_errors(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5023
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    import_response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [None, "+380501234567", None, "Хрещатик", "10"],
        ]),
    )
    file_id = import_response.json()["error_file_id"]

    download_response = await http_client.get(
        url=f"{ERRORS_EXPORT_URL}/{file_id}",
    )

    assert download_response.status_code == status.HTTP_200_OK
    assert download_response.headers["content-type"] == XLSX_CONTENT_TYPE


@pytest.mark.asyncio()
async def test_download_import_errors_not_found(
    http_client: AsyncClient,
) -> None:
    random_uuid = str(uuid.uuid4())
    response = await http_client.get(
        url=f"{ERRORS_EXPORT_URL}/{random_uuid}",
    )

    assert response.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.asyncio()
async def test_download_import_errors_contains_error_column(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    telegram_id = 5024
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    import_response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            [None, "+380501234567", None, "Хрещатик", "10"],
        ]),
    )
    file_id = import_response.json()["error_file_id"]

    download_response = await http_client.get(
        url=f"{ERRORS_EXPORT_URL}/{file_id}",
    )

    wb = load_workbook(BytesIO(download_response.content), read_only=True)
    ws = wb.active
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) >= 1
    last_col_value = data_rows[0][-1]
    assert last_col_value is not None
    assert isinstance(last_col_value, str)


@pytest.mark.asyncio()
async def test_import_skipped_count_matches_rejected(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    telegram_id = 5025
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Валідний", "+380501234567", None, "Хрещатик", "10"],
            [None, "+380502222222", None, "Садова", "5"],
            ["Без вулиці", "+380503333333", None, None, "1"],
        ]),
    )

    assert response.status_code == status.HTTP_200_OK
    assert response.json()["imported"] == 1
    assert response.json()["skipped"] == 2


@pytest.mark.asyncio()
async def test_import_intra_file_duplicate_in_error_report(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    upload_xlsx: Callable[..., dict[str, Any]],
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    telegram_id = 5026
    await setup_full_test_user_with_shop(telegram_id=telegram_id)
    await session.commit()

    headers = customer_headers(telegram_id)

    import_response = await http_client.post(
        url=IMPORT_URL,
        headers=headers,
        files=upload_xlsx([
            ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
            ["Іван Іванов", "+380501234567", None, "Хрещатик", "10"],
        ]),
    )

    data = import_response.json()
    assert data["imported"] == 1
    assert data["skipped"] == 1

    file_id = data["error_file_id"]
    assert file_id is not None

    download_response = await http_client.get(
        url=f"{ERRORS_EXPORT_URL}/{file_id}",
    )

    wb = load_workbook(BytesIO(download_response.content), read_only=True)
    ws = wb.active
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 1
    error_msg = data_rows[0][-1]
    assert "Дубль у файлі" in error_msg


@pytest.mark.asyncio()
async def test_edit_client_propagates_coordinates_to_future_orders(
    http_client: AsyncClient,
    session: AsyncSession,
    customer_headers: Callable[[int], dict[str, Any]],
    setup_full_test_user_with_shop,
    setup_test_client,
    setup_test_order,
) -> None:
    telegram_id = 9001
    _, shop_id = await setup_full_test_user_with_shop(telegram_id=telegram_id)

    old_lat, old_lng = 50.0, 30.0
    new_lat, new_lng = 51.1111, 31.2222

    client_id = await setup_test_client(
        shop_id=shop_id,
        full_name="Клієнт Координати",
        phones=["+380501111111"],
        addresses=[
            {
                "street": "Шевченка",
                "house": "10",
                "latitude": old_lat,
                "longitude": old_lng,
            }
        ],
    )

    today = datetime.now(UTC).date()
    future_date = today + timedelta(days=1)
    past_date = today - timedelta(days=1)

    future_order_id = await setup_test_order(
        shop_id=shop_id,
        client_id=client_id,
        delivery_date=future_date,
        delivery_address={
            "street": "Шевченка",
            "house": "10",
            "coordinates": {"latitude": old_lat, "longitude": old_lng},
        },
    )

    past_order_id = await setup_test_order(
        shop_id=shop_id,
        client_id=client_id,
        delivery_date=past_date,
        delivery_address={
            "street": "Шевченка",
            "house": "10",
            "coordinates": {"latitude": old_lat, "longitude": old_lng},
        },
    )

    await session.commit()

    address_result = await session.execute(
        select(ClientAddress).where(ClientAddress.client_id == client_id)
    )
    address_id = address_result.scalars().first().id

    headers = customer_headers(telegram_id)
    response = await http_client.patch(
        url=f"{BASE_URL}/{client_id}",
        headers=headers,
        json={
            "addresses": [
                {
                    "id": address_id,
                    "street": "Шевченка",
                    "house": "10",
                    "coordinates": {
                        "latitude": new_lat,
                        "longitude": new_lng,
                    },
                    "is_primary": True,
                }
            ]
        },
    )

    assert response.status_code == status.HTTP_200_OK

    await session.flush()

    future_order = await session.get(Order, future_order_id)
    assert future_order.delivery_address.coordinates is not None
    assert future_order.delivery_address.coordinates.latitude == new_lat
    assert future_order.delivery_address.coordinates.longitude == new_lng

    past_order = await session.get(Order, past_order_id)
    assert past_order.delivery_address.coordinates is not None
    assert past_order.delivery_address.coordinates.latitude == old_lat
    assert past_order.delivery_address.coordinates.longitude == old_lng
